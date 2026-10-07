"""Regenerate the attachment fixtures (fpdf2 is not a project dependency):

uv run --with fpdf2 python tests/fixtures/attachments/make_fixtures.py
"""

import io
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from fpdf import FPDF
from pypdf import PdfReader, PdfWriter

HERE = Path(__file__).resolve().parent
MARKER = "ATTACHMENT-MARKER: offline payments via Redsys"
PAGES = [
    "Checkout service specification.\nScope: web checkout, card payments, refunds.",
    f"Payments.\n{MARKER}\nSettlement runs nightly.",
]


def pdf_bytes() -> bytes:
    pdf = FPDF()
    pdf.set_creation_date(datetime(2026, 10, 7, tzinfo=UTC))
    pdf.set_font("helvetica", size=12)
    for text in PAGES:
        pdf.add_page()
        pdf.multi_cell(0, 8, text)
    return bytes(pdf.output())


def encrypted(data: bytes) -> bytes:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
    writer.encrypt(user_password="user-secret", owner_password="owner-secret", algorithm="AES-256")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Checkout service specification", level=1)
    doc.add_paragraph(f"Payments. {MARKER}.")
    table = doc.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, [("Module", "Hours"), ("Payments", "40")], strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def main() -> None:
    pdf = pdf_bytes()
    (HERE / "spec.pdf").write_bytes(pdf)
    (HERE / "encrypted.pdf").write_bytes(encrypted(pdf))
    (HERE / "spec.docx").write_bytes(docx_bytes())
    (HERE / "notes.txt").write_text(
        "Meeting notes\nThe client wants a mobile app.\nLaunch before the summer.\n"
    )


if __name__ == "__main__":
    main()
